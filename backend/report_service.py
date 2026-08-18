"""
Wildlife Population Intelligence System
Report generation service.

Uses the existing WPIS database:
    media_assets
    inference_runs
    image_detections
    audio_predictions
    identified_animals
    species
    users

"""

import os
import re
from datetime import date, datetime, timezone
from uuid import uuid4

from psycopg.types.json import Json

# ============================================================
# PATHS
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

REPORTS_DIR = os.path.join(
    BASE_DIR,
    "uploads",
    "reports",
)

os.makedirs(REPORTS_DIR, exist_ok=True)


# ============================================================
# REPORT TYPES
# ============================================================

REPORT_TYPES = [
    {
        "id": "survey",
        "label": "Wildlife survey report",
        "desc": "Summary of wildlife observations from image and audio monitoring.",
    },
    {
        "id": "population",
        "label": "Species population report",
        "desc": "Species-wise population and observation statistics.",
    },
    {
        "id": "biodiversity",
        "label": "Biodiversity report",
        "desc": "Species richness, diversity and observation distribution.",
    },
    {
        "id": "habitat",
        "label": "Habitat assessment report",
        "desc": "Observation and geographic coverage based on available field metadata.",
    },
    {
        "id": "conservation",
        "label": "Conservation report",
        "desc": "Species population and observation patterns relevant to conservation.",
    },
]


# ============================================================
# REPORT TABLE
# ============================================================

REPORT_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS public.reports (
    id uuid DEFAULT gen_random_uuid() NOT NULL,
    name text NOT NULL,
    type text NOT NULL,
    format text NOT NULL,
    size_bytes bigint NOT NULL,
    file_path text NOT NULL,
    filters jsonb DEFAULT '{}'::jsonb NOT NULL,
    generated_by uuid,
    generated_at timestamp with time zone DEFAULT now() NOT NULL,

    CONSTRAINT reports_pkey PRIMARY KEY (id),

    CONSTRAINT reports_format_check
        CHECK (format IN ('PDF', 'XLSX')),

    CONSTRAINT reports_generated_by_fkey
        FOREIGN KEY (generated_by)
        REFERENCES public.users(id)
        ON DELETE SET NULL
);

CREATE INDEX IF NOT EXISTS idx_reports_generated_at
ON public.reports(generated_at DESC);
"""


def ensure_schema(cur):
    """
    Create only the reports archive table if it does not exist.

    IMPORTANT:
    This deliberately does NOT create monitoring_sites.
    """
    cur.execute(REPORT_TABLE_SQL)


# ============================================================
# HELPERS
# ============================================================

def now_utc():
    return datetime.now(timezone.utc)


def human_size(size):
    if size is None:
        return "0 B"

    size = float(size)

    if size < 1024:
        return f"{int(size)} B"

    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"

    if size < 1024 * 1024 * 1024:
        return f"{size / (1024 * 1024):.1f} MB"

    return f"{size / (1024 * 1024 * 1024):.1f} GB"


def safe_filename(value):
    value = str(value or "report")
    value = re.sub(r"[^a-zA-Z0-9_-]+", "-", value)
    return value.strip("-").lower() or "report"


def parse_date(value):
    if not value:
        return None

    if isinstance(value, date):
        return value

    return date.fromisoformat(str(value))


def date_conditions(from_d=None, to_d=None, alias="ir"):
    """
    Date filtering is based on inference_runs.started_at.
    """
    conditions = []
    params = []

    if from_d:
        conditions.append(
            f"{alias}.started_at >= %s::date"
        )
        params.append(from_d)

    if to_d:
        conditions.append(
            f"{alias}.started_at < (%s::date + INTERVAL '1 day')"
        )
        params.append(to_d)

    return conditions, params


def species_condition(species, alias="s"):
    if not species:
        return [], []

    value = str(species).strip()

    condition = f"""
        (
            LOWER({alias}.name) LIKE LOWER(%s)
            OR LOWER(COALESCE({alias}.scientific_name, '')) LIKE LOWER(%s)
        )
    """

    search = f"%{value}%"

    return [condition], [search, search]


def combined_conditions(from_d=None, to_d=None, species=None):
    conditions = []
    params = []

    date_c, date_p = date_conditions(from_d, to_d)
    conditions.extend(date_c)
    params.extend(date_p)

    species_c, species_p = species_condition(species)
    conditions.extend(species_c)
    params.extend(species_p)

    return conditions, params


def where_clause(conditions):
    if not conditions:
        return ""

    return "WHERE " + " AND ".join(conditions)


# ============================================================
# SURVEY REPORT
# ============================================================

def collect_survey(cur, from_d=None, to_d=None, species=None):
    conditions, params = combined_conditions(
        from_d,
        to_d,
        species,
    )

    where = where_clause(conditions)

    # --------------------------------------------------------
    # Image observations
    # --------------------------------------------------------

    image_query = f"""
        SELECT
            COUNT(*) AS count,
            COUNT(DISTINCT d.species_id) AS species_count,
            COALESCE(AVG(d.confidence), 0) AS avg_confidence
        FROM image_detections d
        JOIN inference_runs ir
            ON ir.id = d.inference_run_id
        LEFT JOIN species s
            ON s.id = d.species_id
        {where}
          AND ir.status = 'completed'
    """ if where else """
        SELECT
            COUNT(*) AS count,
            COUNT(DISTINCT d.species_id) AS species_count,
            COALESCE(AVG(d.confidence), 0) AS avg_confidence
        FROM image_detections d
        JOIN inference_runs ir
            ON ir.id = d.inference_run_id
        LEFT JOIN species s
            ON s.id = d.species_id
        WHERE ir.status = 'completed'
    """

    cur.execute(image_query, params)
    image_stats = cur.fetchone()

    # --------------------------------------------------------
    # Audio observations
    # --------------------------------------------------------

    audio_query = f"""
        SELECT
            COUNT(*) AS count,
            COUNT(DISTINCT ap.species_id) AS species_count,
            COALESCE(AVG(ap.confidence), 0) AS avg_confidence
        FROM audio_predictions ap
        JOIN inference_runs ir
            ON ir.id = ap.inference_run_id
        LEFT JOIN species s
            ON s.id = ap.species_id
        {where}
          AND ir.status = 'completed'
    """ if where else """
        SELECT
            COUNT(*) AS count,
            COUNT(DISTINCT ap.species_id) AS species_count,
            COALESCE(AVG(ap.confidence), 0) AS avg_confidence
        FROM audio_predictions ap
        JOIN inference_runs ir
            ON ir.id = ap.inference_run_id
        LEFT JOIN species s
            ON s.id = ap.species_id
        WHERE ir.status = 'completed'
    """

    cur.execute(audio_query, params)
    audio_stats = cur.fetchone()

    # --------------------------------------------------------
    # Species observations
    # --------------------------------------------------------

    species_query = f"""
        SELECT
            s.name,
            COALESCE(s.scientific_name, '') AS scientific_name,
            COUNT(*) AS observations,
            ROUND(
                AVG(x.confidence)::numeric * 100,
                2
            ) AS avg_confidence
        FROM (
            SELECT
                d.species_id,
                d.confidence,
                d.inference_run_id
            FROM image_detections d

            UNION ALL

            SELECT
                ap.species_id,
                ap.confidence,
                ap.inference_run_id
            FROM audio_predictions ap
        ) x
        JOIN inference_runs ir
            ON ir.id = x.inference_run_id
        JOIN species s
            ON s.id = x.species_id
        {where}
          AND ir.status = 'completed'
        GROUP BY s.id, s.name, s.scientific_name
        ORDER BY observations DESC, s.name
    """ if where else """
        SELECT
            s.name,
            COALESCE(s.scientific_name, '') AS scientific_name,
            COUNT(*) AS observations,
            ROUND(
                AVG(x.confidence)::numeric * 100,
                2
            ) AS avg_confidence
        FROM (
            SELECT
                d.species_id,
                d.confidence,
                d.inference_run_id
            FROM image_detections d

            UNION ALL

            SELECT
                ap.species_id,
                ap.confidence,
                ap.inference_run_id
            FROM audio_predictions ap
        ) x
        JOIN inference_runs ir
            ON ir.id = x.inference_run_id
        JOIN species s
            ON s.id = x.species_id
        WHERE ir.status = 'completed'
        GROUP BY s.id, s.name, s.scientific_name
        ORDER BY observations DESC, s.name
    """

    cur.execute(species_query, params)
    species_rows = cur.fetchall()

    return {
        "title": "Wildlife Survey Report",
        "summary": {
            "Image observations": int(image_stats["count"]),
            "Audio observations": int(audio_stats["count"]),
            "Total observations": (
                int(image_stats["count"]) +
                int(audio_stats["count"])
            ),
            "Species detected": len(species_rows),
            "Average image confidence": round(
                float(image_stats["avg_confidence"]) * 100,
                2,
            ),
            "Average audio confidence": round(
                float(audio_stats["avg_confidence"]) * 100,
                2,
            ),
        },
        "table": [
            {
                "Species": row["name"],
                "Scientific name": row["scientific_name"],
                "Observations": int(row["observations"]),
                "Average confidence": float(row["avg_confidence"] or 0),
            }
            for row in species_rows
        ],
    }


# ============================================================
# POPULATION REPORT
# ============================================================

def collect_population(cur, from_d=None, to_d=None, species=None):
    """
    Population is based on distinct identified_animals.

    Observation counts come from image + audio detections.
    """

    species_conditions = []
    species_params = []

    if species:
        species_conditions, species_params = species_condition(
            species,
            alias="s",
        )

    species_where = where_clause(species_conditions)

    # --------------------------------------------------------
    # Population
    # --------------------------------------------------------

    population_query = f"""
        SELECT
            s.name,
            COALESCE(s.scientific_name, '') AS scientific_name,
            COUNT(DISTINCT ia.id) AS population
        FROM identified_animals ia
        JOIN species s
            ON s.id = ia.species_id
        {species_where}
        GROUP BY s.id, s.name, s.scientific_name
        ORDER BY population DESC, s.name
    """

    cur.execute(population_query, species_params)
    population_rows = cur.fetchall()

    # --------------------------------------------------------
    # Observations and confidence
    # --------------------------------------------------------

    date_conditions_image, date_params_image = date_conditions(
        from_d,
        to_d,
        alias="ir",
    )

    image_conditions = list(date_conditions_image)
    image_params = list(date_params_image)

    if species:
        c, p = species_condition(species, alias="s")
        image_conditions.extend(c)
        image_params.extend(p)

    image_where = where_clause(image_conditions)

    image_query = f"""
        SELECT
            s.name,
            COUNT(*) AS observations,
            COALESCE(AVG(d.confidence), 0) AS confidence
        FROM image_detections d
        JOIN inference_runs ir
            ON ir.id = d.inference_run_id
        JOIN species s
            ON s.id = d.species_id
        {image_where}
        {"AND " if image_where else "WHERE "}ir.status = 'completed'
        GROUP BY s.id, s.name
    """

    cur.execute(image_query, image_params)
    image_rows = cur.fetchall()

    audio_query = f"""
        SELECT
            s.name,
            COUNT(*) AS observations,
            COALESCE(AVG(ap.confidence), 0) AS confidence
        FROM audio_predictions ap
        JOIN inference_runs ir
            ON ir.id = ap.inference_run_id
        JOIN species s
            ON s.id = ap.species_id
        {image_where}
        {"AND " if image_where else "WHERE "}ir.status = 'completed'
        GROUP BY s.id, s.name
    """

    cur.execute(audio_query, image_params)
    audio_rows = cur.fetchall()

    stats = {}

    for row in image_rows:
        name = row["name"]
        stats.setdefault(
            name,
            {
                "observations": 0,
                "confidence_sum": 0,
                "confidence_count": 0,
            },
        )

        stats[name]["observations"] += int(row["observations"])
        stats[name]["confidence_sum"] += (
            float(row["confidence"]) *
            int(row["observations"])
        )
        stats[name]["confidence_count"] += int(row["observations"])

    for row in audio_rows:
        name = row["name"]
        stats.setdefault(
            name,
            {
                "observations": 0,
                "confidence_sum": 0,
                "confidence_count": 0,
            },
        )

        stats[name]["observations"] += int(row["observations"])
        stats[name]["confidence_sum"] += (
            float(row["confidence"]) *
            int(row["observations"])
        )
        stats[name]["confidence_count"] += int(row["observations"])

    table = []

    for row in population_rows:
        name = row["name"]
        stat = stats.get(
            name,
            {
                "observations": 0,
                "confidence_sum": 0,
                "confidence_count": 0,
            },
        )

        avg_confidence = 0

        if stat["confidence_count"]:
            avg_confidence = (
                stat["confidence_sum"] /
                stat["confidence_count"]
            ) * 100

        table.append(
            {
                "Species": name,
                "Scientific name": row["scientific_name"],
                "Population": int(row["population"]),
                "Observations": stat["observations"],
                "Average confidence": round(
                    avg_confidence,
                    2,
                ),
            }
        )

    return {
        "title": "Species Population Report",
        "summary": {
            "Species with identified animals": len(table),
            "Total identified animals": sum(
                row["Population"] for row in table
            ),
            "Total observations": sum(
                row["Observations"] for row in table
            ),
        },
        "table": table,
    }


# ============================================================
# BIODIVERSITY REPORT
# ============================================================

def shannon_index(counts):
    total = sum(counts)

    if total <= 0:
        return 0.0

    result = 0.0

    for count in counts:
        if count <= 0:
            continue

        p = count / total
        result -= p * __import__("math").log(p)

    return result


def collect_biodiversity(cur, from_d=None, to_d=None, species=None):
    conditions, params = combined_conditions(
        from_d,
        to_d,
        species,
    )

    where = where_clause(conditions)

    query = f"""
        SELECT
            s.name,
            COUNT(*) AS observations
        FROM (
            SELECT
                d.species_id,
                d.inference_run_id
            FROM image_detections d

            UNION ALL

            SELECT
                ap.species_id,
                ap.inference_run_id
            FROM audio_predictions ap
        ) x
        JOIN inference_runs ir
            ON ir.id = x.inference_run_id
        JOIN species s
            ON s.id = x.species_id
        {where}
        {"AND " if where else "WHERE "}ir.status = 'completed'
        GROUP BY s.id, s.name
        ORDER BY observations DESC, s.name
    """

    cur.execute(query, params)
    rows = cur.fetchall()

    counts = [int(row["observations"]) for row in rows]

    total = sum(counts)

    table = []

    for row in rows:
        observations = int(row["observations"])

        percentage = (
            observations / total * 100
            if total
            else 0
        )

        table.append(
            {
                "Species": row["name"],
                "Observations": observations,
                "Share (%)": round(percentage, 2),
            }
        )

    return {
        "title": "Biodiversity Report",
        "summary": {
            "Species richness": len(rows),
            "Total observations": total,
            "Shannon diversity index": round(
                shannon_index(counts),
                4,
            ),
        },
        "table": table,
    }


# ============================================================
# HABITAT ASSESSMENT
# ============================================================

def collect_habitat(cur, from_d=None, to_d=None, species=None):
    """
    Your current schema does not contain a habitat table.

    Therefore this report uses the geographic metadata actually
    available in media_assets instead of inventing habitat classes.
    """

    conditions = []
    params = []

    date_c, date_p = date_conditions(
        from_d,
        to_d,
        alias="ir",
    )

    conditions.extend(date_c)
    params.extend(date_p)

    if species:
        c, p = species_condition(
            species,
            alias="s",
        )
        conditions.extend(c)
        params.extend(p)

    where = where_clause(conditions)

    query = f"""
        SELECT
            ma.media_type,
            COUNT(*) AS observations,
            COUNT(
                CASE
                    WHEN ma.latitude IS NOT NULL
                     AND ma.longitude IS NOT NULL
                    THEN 1
                END
            ) AS geotagged
        FROM media_assets ma
        JOIN inference_runs ir
            ON ir.media_id = ma.id
        LEFT JOIN image_detections d
            ON d.inference_run_id = ir.id
        LEFT JOIN audio_predictions ap
            ON ap.inference_run_id = ir.id
        LEFT JOIN species s
            ON s.id = COALESCE(
                d.species_id,
                ap.species_id
            )
        {where}
        {"AND " if where else "WHERE "}ir.status = 'completed'
        GROUP BY ma.media_type
        ORDER BY ma.media_type
    """

    cur.execute(query, params)
    rows = cur.fetchall()

    total_media = 0
    geotagged_media = 0

    table = []

    for row in rows:
        observations = int(row["observations"])
        geotagged = int(row["geotagged"])

        total_media += observations
        geotagged_media += geotagged

        table.append(
            {
                "Media type": row["media_type"],
                "Observations": observations,
                "Geotagged": geotagged,
                "Geotag coverage (%)": round(
                    geotagged / observations * 100,
                    2,
                ) if observations else 0,
            }
        )

    return {
        "title": "Habitat Assessment Report",
        "summary": {
            "Media records analysed": total_media,
            "Geotagged records": geotagged_media,
            "Geographic coverage (%)": round(
                geotagged_media / total_media * 100,
                2,
            ) if total_media else 0,
        },
        "table": table,
    }


# ============================================================
# CONSERVATION REPORT
# ============================================================

def collect_conservation(cur, from_d=None, to_d=None, species=None):
    """
    Conservation observations are based only on actual database
    evidence. No IUCN status is invented here.
    """

    population = collect_population(
        cur,
        from_d,
        to_d,
        species,
    )

    table = []

    for row in population["table"]:
        population_value = row["Population"]
        observations = row["Observations"]

        if population_value == 0:
            observation_ratio = 0
        else:
            observation_ratio = (
                observations / population_value
            )

        table.append(
            {
                "Species": row["Species"],
                "Scientific name": row["Scientific name"],
                "Identified population": population_value,
                "Observations": observations,
                "Average confidence": row["Average confidence"],
                "Observations per identified animal": round(
                    observation_ratio,
                    2,
                ),
            }
        )

    table.sort(
        key=lambda x: x["Observations"],
        reverse=True,
    )

    return {
        "title": "Conservation Observation Report",
        "summary": {
            "Species analysed": len(table),
            "Identified animals": sum(
                row["Identified population"]
                for row in table
            ),
            "Total observations": sum(
                row["Observations"]
                for row in table
            ),
        },
        "table": table,
    }


# ============================================================
# RAW OBSERVATIONS
# ============================================================

def raw_observations(cur, from_d=None, to_d=None, species=None):
    conditions, params = combined_conditions(
        from_d,
        to_d,
        species,
    )

    where = where_clause(conditions)

    query = f"""
        SELECT
            ir.started_at,
            'image' AS source,
            s.name AS species,
            COALESCE(
                s.scientific_name,
                ''
            ) AS scientific_name,
            d.confidence
        FROM image_detections d
        JOIN inference_runs ir
            ON ir.id = d.inference_run_id
        LEFT JOIN species s
            ON s.id = d.species_id
        {where}
        {"AND " if where else "WHERE "}ir.status = 'completed'

        UNION ALL

        SELECT
            ir.started_at,
            'audio' AS source,
            s.name AS species,
            COALESCE(
                s.scientific_name,
                ''
            ) AS scientific_name,
            ap.confidence
        FROM audio_predictions ap
        JOIN inference_runs ir
            ON ir.id = ap.inference_run_id
        LEFT JOIN species s
            ON s.id = ap.species_id
        {where}
        {"AND " if where else "WHERE "}ir.status = 'completed'

        ORDER BY started_at DESC
    """

    # The UNION uses the same parameters twice.
    query_params = params + params

    cur.execute(query, query_params)
    rows = cur.fetchall()

    return [
        {
            "Date": row["started_at"].isoformat()
            if row["started_at"]
            else "",
            "Source": row["source"],
            "Species": row["species"] or "Unknown",
            "Scientific name": row["scientific_name"],
            "Confidence (%)": round(
                float(row["confidence"]) * 100,
                2,
            ),
        }
        for row in rows
    ]


# ============================================================
# PDF
# ============================================================

def render_pdf(
    title,
    filters,
    summary,
    table,
):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import (
        getSampleStyleSheet,
        ParagraphStyle,
    )
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )

    filename = f"{safe_filename(title)}-{uuid4().hex}.pdf"
    path = os.path.join(
        REPORTS_DIR,
        filename,
    )

    doc = SimpleDocTemplate(
        path,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=20,
        leading=24,
        spaceAfter=10,
    )

    normal_style = ParagraphStyle(
        "ReportNormal",
        parent=styles["BodyText"],
        fontSize=9,
        leading=12,
    )

    story = []

    story.append(
        Paragraph(
            "Wildlife Population Intelligence System",
            title_style,
        )
    )

    story.append(
        Paragraph(
            title,
            styles["Heading2"],
        )
    )

    story.append(Spacer(1, 8))

    # --------------------------------------------------------
    # Filters
    # --------------------------------------------------------

    filter_rows = [
        ["From", filters.get("from") or "All"],
        ["To", filters.get("to") or "All"],
        ["Species", filters.get("species") or "All species"],
    ]

    filter_table = Table(
        filter_rows,
        colWidths=[35 * mm, 135 * mm],
    )

    filter_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#eef2e8"),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    9,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(filter_table)
    story.append(Spacer(1, 12))

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Summary",
            styles["Heading3"],
        )
    )

    summary_rows = [
        [str(k), str(v)]
        for k, v in summary.items()
    ]

    if summary_rows:
        summary_table = Table(
            summary_rows,
            colWidths=[85 * mm, 85 * mm],
        )

        summary_table.setStyle(
            TableStyle(
                [
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.4,
                        colors.grey,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (0, -1),
                        "Helvetica-Bold",
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        9,
                    ),
                    (
                        "BACKGROUND",
                        (0, 0),
                        (0, -1),
                        colors.HexColor("#f5f5f5"),
                    ),
                ]
            )
        )

        story.append(summary_table)
        story.append(Spacer(1, 12))

    # --------------------------------------------------------
    # Data table
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Detailed results",
            styles["Heading3"],
        )
    )

    if table:
        columns = list(table[0].keys())

        data = [columns]

        for row in table:
            data.append(
                [
                    str(row.get(column, ""))
                    for column in columns
                ]
            )

        data_table = Table(
            data,
            repeatRows=1,
        )

        data_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#36583b"),
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "FONTNAME",
                        (0, 1),
                        (-1, -1),
                        "Helvetica",
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7.5,
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.35,
                        colors.grey,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "MIDDLE",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#f7f8f5"),
                        ],
                    ),
                ]
            )
        )

        story.append(data_table)

    else:
        story.append(
            Paragraph(
                "No observations matched the selected filters.",
                normal_style,
            )
        )

    story.append(Spacer(1, 15))

    story.append(
        Paragraph(
            f"Generated: {now_utc().strftime('%Y-%m-%d %H:%M UTC')}",
            normal_style,
        )
    )

    story.append(
        Paragraph(
            "Detections are model estimates. Rare or endangered "
            "records should be confirmed before publication.",
            normal_style,
        )
    )

    doc.build(story)

    return path


# ============================================================
# EXCEL
# ============================================================

def render_xlsx(
    title,
    filters,
    summary,
    table,
    raw,
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    filename = f"{safe_filename(title)}-{uuid4().hex}.xlsx"
    path = os.path.join(
        REPORTS_DIR,
        filename,
    )

    workbook = Workbook()

    # --------------------------------------------------------
    # Summary sheet
    # --------------------------------------------------------

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    summary_sheet["A1"] = (
        "Wildlife Population Intelligence System"
    )
    summary_sheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    summary_sheet["A2"] = title
    summary_sheet["A2"].font = Font(
        bold=True,
        size=13,
    )

    row_number = 4

    summary_sheet.cell(
        row=row_number,
        column=1,
        value="Filter",
    )
    summary_sheet.cell(
        row=row_number,
        column=2,
        value="Value",
    )

    row_number += 1

    filter_values = [
        ("From", filters.get("from") or "All"),
        ("To", filters.get("to") or "All"),
        ("Species", filters.get("species") or "All species"),
    ]

    for key, value in filter_values:
        summary_sheet.cell(
            row=row_number,
            column=1,
            value=key,
        )
        summary_sheet.cell(
            row=row_number,
            column=2,
            value=value,
        )
        row_number += 1

    row_number += 1

    summary_sheet.cell(
        row=row_number,
        column=1,
        value="Metric",
    )
    summary_sheet.cell(
        row=row_number,
        column=2,
        value="Value",
    )

    row_number += 1

    for key, value in summary.items():
        summary_sheet.cell(
            row=row_number,
            column=1,
            value=key,
        )
        summary_sheet.cell(
            row=row_number,
            column=2,
            value=value,
        )
        row_number += 1

    # --------------------------------------------------------
    # Results sheet
    # --------------------------------------------------------

    result_sheet = workbook.create_sheet("Results")

    if table:
        columns = list(table[0].keys())

        for column_index, column in enumerate(
            columns,
            start=1,
        ):
            cell = result_sheet.cell(
                row=1,
                column=column_index,
                value=column,
            )
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                "solid",
                fgColor="36583B",
            )
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

        for row_index, row in enumerate(
            table,
            start=2,
        ):
            for column_index, column in enumerate(
                columns,
                start=1,
            ):
                result_sheet.cell(
                    row=row_index,
                    column=column_index,
                    value=row.get(column, ""),
                )

    else:
        result_sheet["A1"] = (
            "No observations matched the selected filters."
        )

    # --------------------------------------------------------
    # Raw observations sheet
    # --------------------------------------------------------

    raw_sheet = workbook.create_sheet(
        "Raw observations"
    )

    if raw:
        columns = list(raw[0].keys())

        for column_index, column in enumerate(
            columns,
            start=1,
        ):
            cell = raw_sheet.cell(
                row=1,
                column=column_index,
                value=column,
            )
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                "solid",
                fgColor="36583B",
            )

        for row_index, row in enumerate(
            raw,
            start=2,
        ):
            for column_index, column in enumerate(
                columns,
                start=1,
            ):
                raw_sheet.cell(
                    row=row_index,
                    column=column_index,
                    value=row.get(column, ""),
                )

    # --------------------------------------------------------
    # Formatting
    # --------------------------------------------------------

    for sheet in workbook.worksheets:
        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                try:
                    value_length = len(str(cell.value))
                    max_length = max(
                        max_length,
                        value_length,
                    )
                except Exception:
                    pass

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                45,
            )

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    workbook.save(path)

    return path


# ============================================================
# GENERATE REPORT
# ============================================================

def generate(
    conn,
    payload,
    generated_by=None,
):
    """
    Generate a report and save it to the reports archive.

    payload:
        {
            "type": "survey",
            "from": "2026-08-01",
            "to": "2026-08-18",
            "species": "Lion",
            "format": "pdf"
        }
    """

    report_type = (
        str(
            payload.get("type") or "survey"
        )
        .strip()
        .lower()
    )

    output_format = (
        str(
            payload.get("format") or "pdf"
        )
        .strip()
        .lower()
    )

    from_d = parse_date(
        payload.get("from")
    )

    to_d = parse_date(
        payload.get("to")
    )

    species = (
        str(payload.get("species")).strip()
        if payload.get("species")
        else None
    )

    if report_type not in {
        item["id"]
        for item in REPORT_TYPES
    }:
        raise ValueError(
            f"Unknown report type: {report_type}"
        )

    if output_format not in {
        "pdf",
        "xlsx",
    }:
        raise ValueError(
            "Format must be pdf or xlsx"
        )

    if from_d and to_d and from_d > to_d:
        raise ValueError(
            "From date cannot be after To date."
        )

    collectors = {
        "survey": collect_survey,
        "population": collect_population,
        "biodiversity": collect_biodiversity,
        "habitat": collect_habitat,
        "conservation": collect_conservation,
    }

    collector = collectors[report_type]

    with conn.cursor() as cur:
        ensure_schema(cur)

        data = collector(
            cur,
            from_d,
            to_d,
            species,
        )

        raw = raw_observations(
            cur,
            from_d,
            to_d,
            species,
        )

        filters = {
            "type": report_type,
            "from": (
                from_d.isoformat()
                if from_d
                else None
            ),
            "to": (
                to_d.isoformat()
                if to_d
                else None
            ),
            "species": species,
        }

        if output_format == "pdf":
            file_path = render_pdf(
                data["title"],
                filters,
                data["summary"],
                data["table"],
            )
            file_format = "PDF"
        else:
            file_path = render_xlsx(
                data["title"],
                filters,
                data["summary"],
                data["table"],
                raw,
            )
            file_format = "XLSX"

        file_size = os.path.getsize(
            file_path
        )

        filename = os.path.basename(
            file_path
        )

        report_name = (
            f"{data['title']} - "
            f"{from_d.isoformat() if from_d else 'all'} "
            f"to "
            f"{to_d.isoformat() if to_d else 'all'}"
        )

        report = cur.execute(
            """
            INSERT INTO reports (
                name,
                type,
                format,
                size_bytes,
                file_path,
                filters,
                generated_by
            )
            VALUES (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )
            RETURNING
                id,
                name,
                type,
                format,
                size_bytes,
                generated_at
            """,
            (
                report_name,
                report_type,
                file_format,
                file_size,
                file_path,
                Json(filters),
                generated_by,
            ),
        ).fetchone()

    conn.commit()

    return {
        "report": {
            "id": str(report["id"]),
            "name": report["name"],
            "type": report["type"],
            "format": report["format"],
            "size": human_size(
                report["size_bytes"]
            ),
            "generatedAt": (
                report["generated_at"].isoformat()
            ),
        },
        "download_url": (
            f"/reports/{report['id']}/download"
        ),
    }


# ============================================================
# REPORT CATALOGUE
# ============================================================

def list_reports(conn):
    with conn.cursor() as cur:
        ensure_schema(cur)

        rows = cur.execute(
            """
            SELECT
                r.id,
                r.name,
                r.type,
                r.format,
                r.size_bytes,
                r.generated_at,
                u.name AS generated_by_name,
                u.email AS generated_by_email
            FROM reports r
            LEFT JOIN users u
                ON u.id = r.generated_by
            ORDER BY r.generated_at DESC
            LIMIT 50
            """
        ).fetchall()

    conn.commit()

    recent = []

    for row in rows:
        by = (
            row["generated_by_name"]
            or row["generated_by_email"]
            or "System"
        )

        recent.append(
            {
                "id": str(row["id"]),
                "name": row["name"],
                "type": row["type"],
                "format": row["format"],
                "size": human_size(
                    row["size_bytes"]
                ),
                "generatedAt": (
                    row["generated_at"].isoformat()
                ),
                "by": by,
            }
        )

    return {
        "types": REPORT_TYPES,
        "recent": recent,
    }


# ============================================================
# GET REPORT FILE
# ============================================================

def get_report(conn, report_id):
    with conn.cursor() as cur:
        row = cur.execute(
            """
            SELECT
                id,
                name,
                type,
                format,
                size_bytes,
                file_path,
                generated_at
            FROM reports
            WHERE id = %s
            """,
            (report_id,),
        ).fetchone()

    return row